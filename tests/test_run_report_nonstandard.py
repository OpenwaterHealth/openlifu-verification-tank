"""Runs with any override are marked NON-STANDARD and cannot write to
the device; standard runs are unmarked."""
from __future__ import annotations

import json
import sys
from dataclasses import asdict
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import run_report  # noqa: E402

from openlifu_verification import ScanConfig  # noqa: E402
from openlifu_verification.report_io import NONSTANDARD_BANNER  # noqa: E402


def _args(tmp_path, *extra):
    return run_report.build_parser().parse_args([
        "--dry-run", "--no-prompt",
        "--prefs", str(tmp_path / "prefs.json"),
        "--reports-root", str(tmp_path / "reports"),
        "--calibration-path", str(tmp_path / "hydrophone.json"),
        *extra,
    ])


def test_standard_args_have_no_reasons(tmp_path):
    assert run_report.nonstandard_reasons(_args(tmp_path), ScanConfig()) == []


@pytest.mark.parametrize("extra, reason", [
    (["--voltage", "25"], "--voltage 25"),
    (["--skip-2d"], "--skip-2d"),
    (["--skip-frequency"], "--skip-frequency"),
    (["--skip-voltage"], "--skip-voltage"),
])
def test_override_options_are_reasons(tmp_path, extra, reason):
    assert run_report.nonstandard_reasons(
        _args(tmp_path, *extra), ScanConfig()) == [reason]


def _run(args):
    seen = {}

    def on_complete(result, write_fn):
        seen["write_fn"] = write_fn

    result = run_report.run(args, on_complete=on_complete)
    rows = []
    if result.files.get("xlsx"):
        from openpyxl import load_workbook
        rows = list(load_workbook(result.files["xlsx"])["Summary"]
                    .iter_rows(values_only=True))
    return result, seen["write_fn"], rows


def test_standard_run_is_unmarked_and_can_write(tmp_path):
    result, write_fn, rows = _run(_args(tmp_path))
    assert result.passed
    assert result.nonstandard_reasons == []
    assert write_fn is not None
    assert ("Configuration", "Standard") in rows
    snapshot = next(result.run_dir.glob("*_scan_config.json"))
    assert json.loads(snapshot.read_text(encoding="utf-8")) == json.loads(
        json.dumps(asdict(ScanConfig())))


def test_nonstandard_run_is_marked_and_cannot_write(tmp_path):
    override = tmp_path / "override.json"
    override.write_text(json.dumps({"scope": {"n_averages": 1}}),
                        encoding="utf-8")

    result, write_fn, rows = _run(
        _args(tmp_path, "--scan-config", str(override)))

    # A passing run, so the only thing blocking the write is the override.
    assert result.passed
    assert write_fn is None
    assert result.nonstandard_reasons == [
        f"--scan-config {override}",
        "scope.n_averages = 1 (standard 3)",
    ]
    snapshot = next(result.run_dir.glob("*_scan_config.json"))
    assert json.loads(snapshot.read_text(encoding="utf-8"))["scope"]["n_averages"] == 1
    assert ("Configuration", NONSTANDARD_BANNER) in rows
    assert ("Non-standard setting", "scope.n_averages = 1 (standard 3)") in rows
